"""
DocxParser: Parse Word files containing CTĐT data.
Extracts data from tables based on header patterns.
"""

import logging
import re

from docx import Document

logger = logging.getLogger(__name__)


class DocxParser:
    """
    Parse a .docx file containing a CTĐT (training program) description.
    Identifies tables by their header row patterns and extracts structured data.
    """

    def __init__(self, file_path):
        self.doc = Document(file_path)
        self.tables = self.doc.tables
        self.warnings = []
        self.errors = []

    def parse_all(self):
        """Parse all known tables and return structured data."""
        result = {
            "general_info": {},
            "plos": [],
            "po_plo_matrix": [],
            "knowledge_blocks": [],
            "courses": [],
            "course_plo_matrix": [],
            "course_descriptions": [],
            "semester_plan": [],
            "performance_indicators": [],
            "assessment_plans": [],
        }

        for i, table in enumerate(self.tables):
            try:
                table_type = self._identify_table(table)
                if table_type == "general_info":
                    result["general_info"] = self._parse_general_info(table)
                elif table_type == "plos":
                    result["plos"] = self._parse_plos(table)
                elif table_type == "po_plo_matrix":
                    result["po_plo_matrix"] = self._parse_po_plo_matrix(table)
                elif table_type == "knowledge_blocks":
                    result["knowledge_blocks"] = self._parse_knowledge_blocks(table)
                elif table_type == "courses":
                    result["courses"] = self._parse_courses(table)
                elif table_type == "course_plo_matrix":
                    result["course_plo_matrix"] = self._parse_course_plo_matrix(table)
                elif table_type == "course_descriptions":
                    result["course_descriptions"] = self._parse_course_descriptions(table)
                elif table_type == "semester_plan":
                    result["semester_plan"] = self._parse_semester_plan(table)
                elif table_type == "pis":
                    result["performance_indicators"] = self._parse_pis(table)
                elif table_type == "assessment_plans":
                    result["assessment_plans"] = self._parse_assessment_plans(table)
            except Exception as e:
                self.errors.append(f"Error parsing table {i}: {e!s}")
                logger.exception(f"Error parsing table {i}")

        return result

    def _identify_table(self, table):
        """Identify table type by inspecting header row content."""
        if not table.rows:
            return None

        # Get first row text
        header_text = " ".join(
            cell.text.strip().lower() for cell in table.rows[0].cells
        )

        # Match patterns
        if "chuẩn đầu ra" in header_text and "mô tả" in header_text:
            return "plos"
        if "mục tiêu" in header_text and ("po" in header_text or "plo" in header_text):
            return "po_plo_matrix"
        if "khối kiến thức" in header_text or "cấu trúc" in header_text:
            return "knowledge_blocks"
        if "mã học phần" in header_text or "mã hp" in header_text:
            if "plo" in header_text or "pi" in header_text:
                return "course_plo_matrix"
            return "courses"
        if "mô tả tóm tắt" in header_text or "mô tả học phần" in header_text:
            return "course_descriptions"
        if "học kỳ" in header_text and ("kế hoạch" in header_text or "đợt" in header_text):
            return "semester_plan"
        if "chỉ số" in header_text and ("pi" in header_text or "đo lường" in header_text):
            return "pis"
        if "đánh giá" in header_text and "plo" in header_text:
            return "assessment_plans"

        # Check for key-value table (general info)
        if len(table.columns) == 2:
            sample_keys = " ".join(
                row.cells[0].text.strip().lower()
                for row in table.rows[:5]
            )
            if any(k in sample_keys for k in ["tên chương trình", "mã chương trình", "trình độ"]):
                return "general_info"

        return None

    def _parse_general_info(self, table):
        """Parse key-value pairs from the general info table."""
        info = {}
        field_map = {
            "tên chương trình": "program_name_vi",
            "tên tiếng anh": "program_name_en",
            "mã chương trình": "program_code",
            "trình độ đào tạo": "education_level",
            "tổng tín chỉ": "total_credits",
            "thời gian đào tạo": "training_duration",
            "tên bằng": "degree_name",
            "số quyết định": "decision_number",
            "mục tiêu chung": "general_objective",
            "điều kiện tuyển sinh": "admission_requirements",
            "điều kiện tốt nghiệp": "graduation_requirements",
            "vị trí việc làm": "career_opportunities",
        }

        for row in table.rows:
            if len(row.cells) >= 2:
                key = row.cells[0].text.strip().lower()
                value = row.cells[1].text.strip()
                for pattern, field in field_map.items():
                    if pattern in key:
                        info[field] = value
                        break

        # Convert total_credits to int
        if "total_credits" in info:
            try:
                info["total_credits"] = int(re.sub(r"\D", "", info["total_credits"]))
            except (ValueError, TypeError):
                self.warnings.append("Could not parse total_credits")

        return info

    def _parse_plos(self, table):
        """Parse PLO table."""
        plos = []
        for i, row in enumerate(table.rows[1:], 1):  # Skip header
            cells = [c.text.strip() for c in row.cells]
            if len(cells) >= 2 and cells[0]:
                plo = {
                    "code": cells[0],
                    "description": cells[1] if len(cells) > 1 else "",
                    "competency_level": cells[2] if len(cells) > 2 else "",
                    "competency_label": cells[3] if len(cells) > 3 else "",
                    "order_index": i,
                }
                plos.append(plo)
        return plos

    def _parse_po_plo_matrix(self, table):
        """Parse PO-PLO mapping matrix."""
        mappings = []
        if not table.rows:
            return mappings

        # First row = header with PO codes
        header_cells = [c.text.strip() for c in table.rows[0].cells]
        po_codes = header_cells[1:]  # Skip first cell (label)

        for row in table.rows[1:]:
            cells = [c.text.strip() for c in row.cells]
            if not cells[0]:
                continue
            plo_code = cells[0]
            for j, po_code in enumerate(po_codes):
                if j + 1 < len(cells) and cells[j + 1].strip():
                    mappings.append({"plo_code": plo_code, "po_code": po_code})

        return mappings

    def _parse_knowledge_blocks(self, table):
        """Parse knowledge block structure."""
        blocks = []
        for i, row in enumerate(table.rows[1:], 1):
            cells = [c.text.strip() for c in row.cells]
            if not cells[0]:
                continue
            block = {
                "name": cells[0],
                "total_credits": self._safe_int(cells[1] if len(cells) > 1 else ""),
                "required_credits": self._safe_int(cells[2] if len(cells) > 2 else ""),
                "elective_credits": self._safe_int(cells[3] if len(cells) > 3 else ""),
                "percentage": cells[4] if len(cells) > 4 else "",
                "order_index": i,
            }
            blocks.append(block)
        return blocks

    def _parse_courses(self, table):
        """Parse course list table."""
        courses = []
        for i, row in enumerate(table.rows[1:], 1):
            cells = [c.text.strip() for c in row.cells]
            if len(cells) < 3 or not cells[0]:
                continue
            course = {
                "order_number": cells[0] if len(cells) > 0 else "",
                "code": cells[1] if len(cells) > 1 else "",
                "name_vi": cells[2] if len(cells) > 2 else "",
                "total_credits": self._safe_int(cells[3] if len(cells) > 3 else ""),
                "theory_credits": self._safe_int(cells[4] if len(cells) > 4 else ""),
                "practice_credits": self._safe_int(cells[5] if len(cells) > 5 else ""),
                "is_required": True,  # Default
                "semester": self._safe_int(cells[7] if len(cells) > 7 else ""),
            }
            courses.append(course)
        return courses

    def _parse_course_plo_matrix(self, table):
        """Parse course-PLO contribution matrix."""
        contributions = []
        if not table.rows:
            return contributions

        header_cells = [c.text.strip() for c in table.rows[0].cells]
        pi_codes = header_cells[2:]  # Skip order and course code columns

        for row in table.rows[1:]:
            cells = [c.text.strip() for c in row.cells]
            if len(cells) < 3 or not cells[1]:
                continue
            course_code = cells[1]
            for j, pi_code in enumerate(pi_codes):
                idx = j + 2
                if idx < len(cells) and cells[idx].strip():
                    level = self._safe_int(cells[idx])
                    if level > 0:
                        contributions.append({
                            "course_code": course_code,
                            "pi_code": pi_code,
                            "contribution_level": level,
                        })

        return contributions

    def _parse_course_descriptions(self, table):
        """Parse course description table."""
        descriptions = []
        for row in table.rows[1:]:
            cells = [c.text.strip() for c in row.cells]
            if len(cells) >= 2 and cells[0]:
                descriptions.append({
                    "code": cells[0],
                    "description": cells[1] if len(cells) > 1 else "",
                })
        return descriptions

    def _parse_semester_plan(self, table):
        """Parse semester teaching plan."""
        plans = []
        current_semester = None
        for i, row in enumerate(table.rows[1:], 1):
            cells = [c.text.strip() for c in row.cells]
            if not cells:
                continue

            # Detect semester header rows
            full_text = " ".join(cells).lower()
            sem_match = re.search(r"học kỳ\s*(\d+)", full_text)
            if sem_match:
                current_semester = int(sem_match.group(1))
                continue

            # Regular course row
            course_code = None
            for cell in cells:
                if re.match(r"^[A-Z]{2,5}\d{3,4}$", cell):
                    course_code = cell
                    break

            if course_code and current_semester:
                plans.append({
                    "semester_number": current_semester,
                    "course_code": course_code,
                    "order_index": i,
                })
        return plans

    def _parse_pis(self, table):
        """Parse performance indicator table."""
        pis = []
        for i, row in enumerate(table.rows[1:], 1):
            cells = [c.text.strip() for c in row.cells]
            if len(cells) >= 2 and cells[0]:
                pis.append({
                    "plo_code": cells[0] if "." not in cells[0] else cells[0].split(".")[0],
                    "code": cells[0],
                    "description": cells[1] if len(cells) > 1 else "",
                    "order_index": i,
                })
        return pis

    def _parse_assessment_plans(self, table):
        """Parse PLO assessment plan table."""
        plans = []
        for row in table.rows[1:]:
            cells = [c.text.strip() for c in row.cells]
            if len(cells) < 3 or not cells[0]:
                continue
            plan = {
                "pi_code": cells[0],
                "contributing_courses_text": cells[1] if len(cells) > 1 else "",
                "sample_course_code": cells[2] if len(cells) > 2 else "",
                "direct_evidence": cells[3] if len(cells) > 3 else "",
                "assessment_tool": cells[4] if len(cells) > 4 else "",
                "expected_standard": cells[5] if len(cells) > 5 else "",
                "assessment_schedule": cells[6] if len(cells) > 6 else "",
            }
            plans.append(plan)
        return plans

    @staticmethod
    def _safe_int(value):
        """Convert value to int, return 0 if not possible."""
        try:
            return int(re.sub(r"\D", "", str(value)))
        except (ValueError, TypeError):
            return 0
